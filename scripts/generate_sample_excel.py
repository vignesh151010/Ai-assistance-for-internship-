import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def generate_sample_internship_sheet(filepath: str = "sample_interns_tasks.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Interns & Tasks"

    # Define headers
    headers = [
        "Candidate Name",
        "Candidate Email",
        "Internship Role",
        "Assigned Task",
        "Task Description",
        "Task Requirements",
        "Expected Deliverables",
        "Required Skills",
        "Deadline",
        "Internship Start Date",
        "Internship End Date",
        "Assigned Mentor"
    ]

    ws.append(headers)

    # Style header row
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")  # Indigo
    thin_border = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    sample_rows = [
        [
            "Rohan Verma",
            "rohan.verma@example.com",
            "Frontend Engineer Intern",
            "Responsive Developer Portfolio Platform",
            "Design and build an accessible, highly-aesthetic developer portfolio displaying project showcases, skills matrix, and interactive contact modal.",
            "Must be responsive across mobile/desktop; CSS Grid/Flexbox; Interactive project filtering; Form validation with visual feedback; Clean GitHub repo.",
            "Production bundle, source files, automated test verification, README with deployment notes.",
            "React, CSS3, JavaScript, HTML5, Git",
            "2026-11-15",
            "2026-10-01",
            "2026-12-31",
            "mentor@techcorp.com"
        ],
        [
            "Priya Sundaram",
            "priya.sundaram@example.com",
            "Backend Systems Intern",
            "Distributed Event Ingestion & Rate Limiting Microservice",
            "Implement high-throughput API gateway middleware supporting sliding-window rate limiting, token bucket algorithms, and Redis-backed state.",
            "P99 latency under 15ms; Unit tests > 85% coverage; Docker containerization; Comprehensive OpenAPI specification.",
            "FastAPI microservice code, Dockerfile, benchmark test suite report, architecture diagram.",
            "Python, FastAPI, Redis, Docker, Pytest",
            "2026-11-20",
            "2026-10-01",
            "2026-12-31",
            "mentor@techcorp.com"
        ],
        [
            "Ananya Sen",
            "ananya.sen@example.com",
            "AI & Data Science Intern",
            "Automated Document Summarization & Sentiment Pipeline",
            "Develop an NLP pipeline capable of extracting key entities, generating concise abstracts, and evaluating sentiment from enterprise PDF reports.",
            "Multi-format document ingestion (PDF, TXT); Entity extraction with confidence scores; Unit tests covering edge cases.",
            "Python pipeline module, Jupyter demo notebook, evaluation benchmark summary.",
            "Python, Transformers, PyTorch, Pandas, Scikit-learn",
            "2026-11-25",
            "2026-10-01",
            "2026-12-31",
            "mentor@techcorp.com"
        ],
        [
            "Vikram Malhotra",
            "vikram.malhotra@example.com",
            "Cloud DevOps Intern",
            "Automated CI/CD Pipeline & Kubernetes Deployment Helm Chart",
            "Construct robust GitHub Actions workflow automating linting, security scans, unit tests, Docker container builds, and deployment to local Minikube cluster.",
            "Zero plaintext secrets; Multi-stage Docker builds; Helm chart with configurable values.yaml; Rollback strategy.",
            "GitHub Actions workflows, Helm chart directory, documentation and deployment verification log.",
            "Docker, Kubernetes, Helm, GitHub Actions, Linux",
            "2026-11-30",
            "2026-10-01",
            "2026-12-31",
            "mentor@techcorp.com"
        ]
    ]

    for row_idx, r in enumerate(sample_rows, start=2):
        ws.append(r)
        row_fill = PatternFill(start_color="F9FAFB" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for cell in ws[row_idx]:
            cell.alignment = Alignment(vertical="center")
            cell.border = thin_border
            cell.fill = row_fill

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 40)

    ws.row_dimensions[1].height = 28

    wb.save(filepath)
    print(f"Sample spreadsheet successfully generated at: {os.path.abspath(filepath)}")

if __name__ == "__main__":
    generate_sample_internship_sheet()
