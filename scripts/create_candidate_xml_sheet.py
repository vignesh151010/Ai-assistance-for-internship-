import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATA = [
    {
        "name": "Rohan Verma",
        "email": "rohan.verma@example.com",
        "role": "Frontend Engineer Intern",
        "task": "Responsive Developer Portfolio Platform",
        "desc": "Design and build an accessible, highly-aesthetic developer portfolio displaying project showcases, skills matrix, and interactive contact modal.",
        "reqs": "Must be responsive across mobile/desktop; CSS Grid/Flexbox; Interactive project filtering; Form validation with visual feedback; Clean GitHub repo.",
        "deliv": "Production bundle, source files, automated test verification, README with deployment notes.",
        "skills": "React, CSS3, JavaScript, HTML5, Git",
        "deadline": "2026-11-15",
        "start": "2026-10-01",
        "end": "2026-12-31",
        "mentor": "mentor@techcorp.com"
    },
    {
        "name": "Priya Sundaram",
        "email": "priya.sundaram@example.com",
        "role": "Backend Systems Intern",
        "task": "Distributed Event Ingestion & Rate Limiting Microservice",
        "desc": "Implement high-throughput API gateway middleware supporting sliding-window rate limiting, token bucket algorithms, and Redis-backed state.",
        "reqs": "P99 latency under 15ms; Unit tests > 85% coverage; Docker containerization; Comprehensive OpenAPI specification.",
        "deliv": "FastAPI microservice code, Dockerfile, benchmark test suite report, architecture diagram.",
        "skills": "Python, FastAPI, Redis, Docker, Pytest",
        "deadline": "2026-11-20",
        "start": "2026-10-01",
        "end": "2026-12-31",
        "mentor": "mentor@techcorp.com"
    },
    {
        "name": "Ananya Sen",
        "email": "ananya.sen@example.com",
        "role": "AI & Data Science Intern",
        "task": "Automated Document Summarization & Sentiment Pipeline",
        "desc": "Develop an NLP pipeline capable of extracting key entities, generating concise abstracts, and evaluating sentiment from enterprise PDF reports.",
        "reqs": "Multi-format document ingestion (PDF, TXT); Entity extraction with confidence scores; Unit tests covering edge cases.",
        "deliv": "Python pipeline module, Jupyter demo notebook, evaluation benchmark summary.",
        "skills": "Python, Transformers, PyTorch, Pandas, Scikit-learn",
        "deadline": "2026-11-25",
        "start": "2026-10-01",
        "end": "2026-12-31",
        "mentor": "mentor@techcorp.com"
    },
    {
        "name": "Vikram Malhotra",
        "email": "vikram.malhotra@example.com",
        "role": "Cloud DevOps Intern",
        "task": "Automated CI/CD Pipeline & Kubernetes Deployment Helm Chart",
        "desc": "Construct robust GitHub Actions workflow automating linting, security scans, unit tests, Docker container builds, and deployment to local Minikube cluster.",
        "reqs": "Zero plaintext secrets; Multi-stage Docker builds; Helm chart with configurable values.yaml; Rollback strategy.",
        "deliv": "GitHub Actions workflows, Helm chart directory, documentation and deployment verification log.",
        "skills": "Docker, Kubernetes, Helm, GitHub Actions, Linux",
        "deadline": "2026-11-30",
        "start": "2026-10-01",
        "end": "2026-12-31",
        "mentor": "mentor@techcorp.com"
    },
    {
        "name": "Sneha Kulkarni",
        "email": "sneha.kulkarni@example.com",
        "role": "Cybersecurity & QA Intern",
        "task": "Automated API Vulnerability & OWASP ZAP Security Scanner",
        "desc": "Build automated vulnerability test suite integrating OWASP Top 10 security checks, input fuzzing, and JWT tampering verification against REST APIs.",
        "reqs": "Detect XSS, SQLi, and IDOR vulnerabilities; Automated vulnerability scoring; HTML and JSON audit report output.",
        "deliv": "Security test scripts, OWASP configuration profile, vulnerability report sample.",
        "skills": "Python, OWASP ZAP, Security Testing, Burp Suite, Bash",
        "deadline": "2026-12-05",
        "start": "2026-10-01",
        "end": "2026-12-31",
        "mentor": "mentor@techcorp.com"
    }
]

def generate_excel_xml(filepath: str = "candidate_task_assignment.xml"):
    """Generate Microsoft Excel XML (SpreadsheetML) format"""
    xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<?mso-application progid="Excel.Sheet"?>
<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:o="urn:schemas-microsoft-com:office:office"
 xmlns:x="urn:schemas-microsoft-com:office:excel"
 xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"
 xmlns:html="http://www.w3.org/TR/REC-html40">
 <DocumentProperties xmlns="urn:schemas-microsoft-com:office:office">
  <Author>TechMentorship System</Author>
  <Title>Candidate Task Assignment Sheet</Title>
  <Created>2026-10-26T12:00:00Z</Created>
 </DocumentProperties>
 <Styles>
  <Style ss:ID="Default" ss:Name="Normal">
   <Alignment ss:Vertical="Center"/>
   <Borders/>
   <Font ss:FontName="Calibri" ss:Size="11" ss:Color="#000000"/>
   <Interior/>
   <NumberFormat/>
   <Protection/>
  </Style>
  <Style ss:ID="Header">
   <Alignment ss:Horizontal="Center" ss:Vertical="Center" ss:WrapText="1"/>
   <Borders>
    <Border ss:Position="Bottom" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#4338CA"/>
    <Border ss:Position="Left" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#4338CA"/>
    <Border ss:Position="Right" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#4338CA"/>
    <Border ss:Position="Top" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#4338CA"/>
   </Borders>
   <Font ss:FontName="Calibri" ss:Size="11" ss:Color="#FFFFFF" ss:Bold="1"/>
   <Interior ss:Color="#312E81" ss:Pattern="Solid"/>
  </Style>
  <Style ss:ID="RowEven">
   <Alignment ss:Vertical="Center" ss:WrapText="1"/>
   <Borders>
    <Border ss:Position="Bottom" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Left" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Right" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Top" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
   </Borders>
   <Font ss:FontName="Calibri" ss:Size="10" ss:Color="#1E293B"/>
   <Interior ss:Color="#F8FAFC" ss:Pattern="Solid"/>
  </Style>
  <Style ss:ID="RowOdd">
   <Alignment ss:Vertical="Center" ss:WrapText="1"/>
   <Borders>
    <Border ss:Position="Bottom" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Left" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Right" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
    <Border ss:Position="Top" ss:LineStyle="Continuous" ss:Weight="1" ss:Color="#E2E8F0"/>
   </Borders>
   <Font ss:FontName="Calibri" ss:Size="10" ss:Color="#1E293B"/>
   <Interior ss:Color="#FFFFFF" ss:Pattern="Solid"/>
  </Style>
 </Styles>
 <Worksheet ss:Name="Candidate Task Assignments">
  <Table ss:DefaultRowHeight="20">
   <Column ss:Width="120"/>
   <Column ss:Width="160"/>
   <Column ss:Width="150"/>
   <Column ss:Width="200"/>
   <Column ss:Width="260"/>
   <Column ss:Width="260"/>
   <Column ss:Width="220"/>
   <Column ss:Width="150"/>
   <Column ss:Width="90"/>
   <Column ss:Width="90"/>
   <Column ss:Width="90"/>
   <Column ss:Width="160"/>
   <Row ss:Height="28" ss:StyleID="Header">
    <Cell><Data ss:Type="String">Candidate Name</Data></Cell>
    <Cell><Data ss:Type="String">Candidate Email</Data></Cell>
    <Cell><Data ss:Type="String">Internship Role</Data></Cell>
    <Cell><Data ss:Type="String">Assigned Task</Data></Cell>
    <Cell><Data ss:Type="String">Task Description</Data></Cell>
    <Cell><Data ss:Type="String">Task Requirements</Data></Cell>
    <Cell><Data ss:Type="String">Expected Deliverables</Data></Cell>
    <Cell><Data ss:Type="String">Required Skills</Data></Cell>
    <Cell><Data ss:Type="String">Deadline</Data></Cell>
    <Cell><Data ss:Type="String">Internship Start Date</Data></Cell>
    <Cell><Data ss:Type="String">Internship End Date</Data></Cell>
    <Cell><Data ss:Type="String">Assigned Mentor</Data></Cell>
   </Row>
"""
    for idx, d in enumerate(DATA):
        style = "RowEven" if idx % 2 == 0 else "RowOdd"
        xml_content += f"""   <Row ss:Height="40" ss:StyleID="{style}">
    <Cell><Data ss:Type="String">{d['name']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['email']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['role']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['task']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['desc']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['reqs']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['deliv']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['skills']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['deadline']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['start']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['end']}</Data></Cell>
    <Cell><Data ss:Type="String">{d['mentor']}</Data></Cell>
   </Row>
"""

    xml_content += """  </Table>
 </Worksheet>
</Workbook>"""

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(xml_content)
    print(f"Generated XML Spreadsheet at: {os.path.abspath(filepath)}")

def generate_xlsx(filepath: str = "candidate_task_assignment.xlsx"):
    """Generate modern .xlsx format"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidate Task Assignment"

    headers = [
        "Candidate Name", "Candidate Email", "Internship Role", "Assigned Task",
        "Task Description", "Task Requirements", "Expected Deliverables",
        "Required Skills", "Deadline", "Internship Start Date", "Internship End Date",
        "Assigned Mentor"
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row_idx, d in enumerate(DATA, start=2):
        row_vals = [
            d["name"], d["email"], d["role"], d["task"], d["desc"],
            d["reqs"], d["deliv"], d["skills"], d["deadline"], d["start"],
            d["end"], d["mentor"]
        ]
        ws.append(row_vals)
        fill = PatternFill(start_color="F8FAFC" if row_idx % 2 == 0 else "FFFFFF", fill_type="solid")
        for cell in ws[row_idx]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = thin_border
            cell.fill = fill

    for col in ws.columns:
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = 24

    ws.row_dimensions[1].height = 28
    wb.save(filepath)
    print(f"Generated XLSX Spreadsheet at: {os.path.abspath(filepath)}")

if __name__ == "__main__":
    generate_excel_xml()
    generate_xlsx()
